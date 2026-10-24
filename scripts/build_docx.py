#!/usr/bin/env python3
"""Build an Arabic (RTL) Word document from one or more Markdown files.

Pipeline: Markdown --pandoc--> .docx --python-docx post-processing--> final .docx

The post-processing step is what makes the document look like a proper Arabic
thesis chapter: right-to-left paragraphs and tables, Arabic-friendly fonts,
an optional cover page, a table of contents field and page numbers.

Usage:
  python3 scripts/build_docx.py OUT.docx IN1.md [IN2.md ...] \
      [--title "..."] [--title-en "..."] [--subtitle "..."] [--author "..."] \
      [--supervisor "..."] [--university "..."] [--program "..."] [--year 2026] \
      [--no-toc] [--no-cover] [--body-font "Simplified Arabic"] [--latin-font "Times New Roman"]

Only pandoc and python-docx (both preinstalled) are required.
"""
from __future__ import annotations

import argparse
import copy
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

LTR_STYLES = {"Source Code", "Verbatim Char", "SourceCode"}


# --------------------------------------------------------------------------- pandoc
def run_pandoc(md_files: list[Path], out: Path, toc: bool, toc_title: str) -> None:
    cmd = [
        "pandoc", *map(str, md_files),
        "-f", "markdown+pipe_tables+grid_tables+multiline_tables+implicit_figures+footnotes+fenced_divs+bracketed_spans",
        "-t", "docx", "-o", str(out),
        "-M", "lang=ar", "-M", "dir=rtl",
        "--resource-path", ":".join({str(p.parent) for p in md_files} | {"."}),
    ]
    if toc:
        cmd += ["--toc", "--toc-depth=3", "-M", f"toc-title={toc_title}"]
    subprocess.run(cmd, check=True)


# --------------------------------------------------------------------------- XML helpers
def _set_rfonts(rpr, latin: str, arabic: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:cs"), arabic)
    rfonts.set(qn("w:eastAsia"), latin)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def _ensure(parent, tag: str):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el


def _set_bidi(paragraph, rtl: bool = True) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if rtl:
        if bidi is None:
            bidi = OxmlElement("w:bidi")
            ppr.append(bidi)
        bidi.set(qn("w:val"), "1")
    elif bidi is not None:
        ppr.remove(bidi)


def _style_fonts(doc: Document, latin: str, arabic: str, body_pt: int) -> None:
    """Apply fonts + sizes on the styles themselves so every run inherits them."""
    heading_sizes = {"Heading 1": 20, "Heading 2": 16, "Heading 3": 14, "Heading 4": 13, "Title": 26, "Subtitle": 16}
    for style in doc.styles:
        try:
            el = style.element
        except AttributeError:
            continue
        rpr = el.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            el.append(rpr)
        name = style.name or ""
        if name in LTR_STYLES:
            _set_rfonts(rpr, "Consolas", "Consolas")
            continue
        _set_rfonts(rpr, latin, arabic)
        size = heading_sizes.get(name)
        if name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
            size = body_pt
        if size:
            for tag in ("w:sz", "w:szCs"):
                e = _ensure(rpr, tag)
                e.set(qn("w:val"), str(size * 2))
        if name.startswith("Heading") or name in ("Title", "Subtitle"):
            color = _ensure(rpr, "w:color")
            color.set(qn("w:val"), "1F3864")
            b = _ensure(rpr, "w:b"); b.set(qn("w:val"), "1")
            bcs = _ensure(rpr, "w:bCs"); bcs.set(qn("w:val"), "1")
        # Paragraph-level RTL on styles too
        ppr = el.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            el.insert(0, ppr) if el.find(qn("w:rPr")) is None else el.insert(list(el).index(rpr), ppr)
        if name not in LTR_STYLES:
            bidi = _ensure(ppr, "w:bidi"); bidi.set(qn("w:val"), "1")
            if name in ("Normal", "Body Text", "First Paragraph"):
                jc = _ensure(ppr, "w:jc"); jc.set(qn("w:val"), "both")
                sp = _ensure(ppr, "w:spacing"); sp.set(qn("w:line"), "360"); sp.set(qn("w:lineRule"), "auto"); sp.set(qn("w:after"), "120")


def _rtl_tables(doc: Document) -> None:
    for table in doc.tables:
        tblpr = table._tbl.tblPr
        bv = _ensure(tblpr, "w:bidiVisual"); bv.set(qn("w:val"), "1")
        jc = _ensure(tblpr, "w:jc"); jc.set(qn("w:val"), "center")
        borders = tblpr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                b = OxmlElement(f"w:{edge}")
                b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:space"), "0"); b.set(qn("w:color"), "808080")
                borders.append(b)
            tblpr.append(borders)
        # header row shading + bold
        if table.rows:
            for cell in table.rows[0].cells:
                tcpr = cell._tc.get_or_add_tcPr()
                shd = _ensure(tcpr, "w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "D9E2F3")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _set_bidi(p, True)
                    if p.alignment is None:
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _rtl_paragraphs(doc: Document) -> None:
    for p in doc.paragraphs:
        sname = p.style.name if p.style is not None else ""
        if sname in LTR_STYLES:
            _set_bidi(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue
        _set_bidi(p, True)
        # Center figure captions / images
        if sname in ("Caption", "Image Caption", "Table Caption", "Figure"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _page_setup(doc: Document) -> None:
    for s in doc.sections:
        s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5); s.right_margin = Cm(3.0)
        sectpr = s._sectPr
        bidi = _ensure(sectpr, "w:bidi")  # RTL section (gutter on the right)
        bidi.set(qn("w:val"), "1")
        # footer page number
        footer = s.footer
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not fp.runs:
            run = fp.add_run()
            for tag, text in (("begin", None), (None, " PAGE "), ("end", None)):
                if tag:
                    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
                else:
                    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)


def _insert_paragraph_before(anchor_el, parent, text: str = "", style: str | None = None, size: int | None = None,
                             bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER, color: str | None = None):
    """Insert a new paragraph before an arbitrary body element (paragraph, table or TOC sdt)."""
    new_p = OxmlElement("w:p")
    anchor_el.addprevious(new_p)
    from docx.text.paragraph import Paragraph
    para = Paragraph(new_p, parent)
    if style:
        para.style = style
    para.alignment = align
    _set_bidi(para, True)
    if text:
        run = para.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
            run._r.get_or_add_rPr()
            szcs = _ensure(run._r.rPr, "w:szCs"); szcs.set(qn("w:val"), str(size * 2))
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return para


def _add_cover(doc: Document, a: argparse.Namespace) -> None:
    body = doc.element.body
    # first real body child (pandoc puts the TOC in a <w:sdt> before the first paragraph)
    first_el = next((el for el in body.iterchildren() if el.tag != qn("w:sectPr")), None)
    if first_el is None:
        first_el = doc.add_paragraph()._p
    parent = doc._body
    lines = [
        (a.university, 16, True, "1F3864"),
        (a.faculty, 13, False, None),
        (a.program, 13, False, None),
        ("", 10, False, None), ("", 10, False, None),
        (a.title, 24, True, "1F3864"),
        (a.title_en, 16, True, "1F3864"),
        (a.subtitle, 14, False, "595959"),
        ("", 10, False, None), ("", 10, False, None),
        (f"إعداد الطالب: {a.author}" if a.author else "", 14, False, None),
        (f"إشراف: {a.supervisor}" if a.supervisor else "", 14, False, None),
        ("", 10, False, None),
        (str(a.year) if a.year else "", 14, False, None),
    ]
    for text, size, bold, color in lines:
        if text is None:
            continue
        _insert_paragraph_before(first_el, parent, text, size=size, bold=bold, color=color)
    # page break after cover
    br = _insert_paragraph_before(first_el, parent, "")
    br.add_run().add_break(WD_BREAK.PAGE)


def _mark_toc_dirty(doc: Document) -> None:
    """Ask Word to refresh the TOC on open (pandoc writes an empty TOC field)."""
    settings = doc.settings.element
    uf = settings.find(qn("w:updateFields"))
    if uf is None:
        uf = OxmlElement("w:updateFields")
        settings.append(uf)
    uf.set(qn("w:val"), "true")


# --------------------------------------------------------------------------- main
def build(args: argparse.Namespace) -> Path:
    md_files = [Path(p) for p in args.inputs]
    for p in md_files:
        if not p.exists():
            sys.exit(f"missing input: {p}")
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / "pandoc.docx"
        run_pandoc(md_files, tmp, toc=not args.no_toc, toc_title=args.toc_title)
        doc = Document(str(tmp))
        _style_fonts(doc, args.latin_font, args.body_font, args.body_pt)
        _rtl_paragraphs(doc)
        _rtl_tables(doc)
        _page_setup(doc)
        if not args.no_cover and args.title:
            _add_cover(doc, args)
        if not args.no_toc:
            _mark_toc_dirty(doc)
        core = doc.core_properties
        core.title = args.title or out.stem
        core.author = args.author or ""
        core.language = "ar-SY"
        doc.save(str(out))
    if not args.no_toc and not args.no_update_toc:
        # Write real page numbers into the TOC via headless LibreOffice (see lo_update_toc.py).
        helper = Path(__file__).with_name("lo_update_toc.py")
        try:
            subprocess.run([sys.executable, str(helper), str(out)], check=True, timeout=600)
        except Exception as exc:  # the document is still valid; Word refreshes the TOC on open
            print(f"warning: TOC page numbers not pre-rendered ({exc}); Word will update the field on open")
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("output")
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("--title", default="")
    ap.add_argument("--title-en", default="")
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--author", default="كنان محمد قرحالي")
    ap.add_argument("--supervisor", default="د. باسل الخطيب")
    ap.add_argument("--university", default="الجامعة الافتراضية السورية")
    ap.add_argument("--faculty", default="")
    ap.add_argument("--program", default="ماجستير التأهيل والتخصص في علوم الوب (MWS)")
    ap.add_argument("--year", default="2026")
    ap.add_argument("--toc-title", default="فهرس المحتويات")
    ap.add_argument("--no-toc", action="store_true")
    ap.add_argument("--no-update-toc", action="store_true", help="skip the LibreOffice pass that fills TOC page numbers")
    ap.add_argument("--no-cover", action="store_true")
    ap.add_argument("--body-font", default="Simplified Arabic")
    ap.add_argument("--latin-font", default="Times New Roman")
    ap.add_argument("--body-pt", type=int, default=14)
    args = ap.parse_args()
    out = build(args)
    print(f"built {out} ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
