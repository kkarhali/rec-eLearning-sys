#!/usr/bin/env python3
"""مولّد العرض التقديمي (python-pptx فقط، لا npm ولا CDN).

الاستعمال:
    python3 scripts/build_pptx.py                      # يقرأ docs/working/07-presentation.yaml أو .json
    python3 scripts/build_pptx.py --spec X.json --out Y.pptx

يقرأ مخطط الشرائح ويُخرج pptx عربياً RTL بنسبة 16:9. أنواع الشرائح:
title, section, bullets, cards, stats, flow, timeline, two-column, image,
table, chart, quote, qa, placeholder. التفاصيل في docs/working/07-presentation-README.md
"""
from __future__ import annotations

import argparse
import copy
import json
import math
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
if not (ROOT / "docs/working").exists():      # تشغيل من نسخة خارج المستودع
    ROOT = Path.cwd()

# ---------------------------------------------------------------- الهوية البصرية
# رموز التصميم من docs/working/03-analysis-design.md §9.1 (الجدول 16)
PRIMARY = "0F766E"
PRIMARY_SOFT = "CCFBF1"
PRIMARY_DARK = "134E4A"      # النص الأساسي وخلفية العنوان/الخاتمة
PRIMARY_DEEP = "115E59"
ACCENT = "B45309"
ACCENT_SOFT = "FEF3C7"
ACCENT_INK = "92400E"
BG = "F8FAFC"
CARD = "FFFFFF"
MUTED = "475569"
BORDER = "CBD5E1"
WHITE = "FFFFFF"

FONT = "Arial"               # خط نظام؛ يحلّ محلّه خط عربي احتياطي عند غيابه
LANG = "ar-SA"
LRM = "‎"

W_IN, H_IN = 13.333, 7.5
MX = 0.6                     # الهامش الجانبي
CONTENT_W = W_IN - 2 * MX
RIGHT = W_IN - MX
BOTTOM = 6.75


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- نص RTL
def _set_cs(rpr, face: str) -> None:
    """يضيف a:cs (الخط للنص المعقّد/العربي) بعد a:latin بالترتيب الصحيح."""
    for tag in ("a:ea", "a:cs"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    latin = rpr.find(qn("a:latin"))
    cs = etree.SubElement(rpr, qn("a:cs"))
    cs.set("typeface", face)
    if latin is not None:
        latin.addnext(cs)


def style_run(r, size, color, bold=False, italic=False):
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = FONT
    rpr = r._r.get_or_add_rPr()
    rpr.set("lang", LANG)
    _set_cs(rpr, FONT)


def rtl_par(p, align="r"):
    """a:pPr بسمة rtl="1" و algn."""
    ppr = p._p.get_or_add_pPr()
    ppr.set("rtl", "1")
    ppr.set("algn", align)


def add_runs(p, text, size, color, bold=False, italic=False, bold_color=None):
    """يدعم **غامق** داخل النص."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_b = part.startswith("**") and part.endswith("**") and len(part) > 4
        r = p.add_run()
        r.text = part[2:-2] if is_b else part
        style_run(r, size, (bold_color or color) if is_b else color, bold or is_b, italic)


def set_par_spacing(p, before=0, after=0, line=None):
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if line:
        p.line_spacing = line


def textbox(slide, x, y, w, h, paras, size=18, color=PRIMARY_DARK, bold=False,
            italic=False, align="r", anchor="t", name=None, margin=0.05,
            line=1.15, after=0):
    """paras: نص أو قائمة نصوص (فقرات). كل فقرة RTL."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        rtl_par(p, align)
        set_par_spacing(p, 0, after, line)
        add_runs(p, t, size, color, bold, italic, bold_color=None)
    return tb


def shape_text(shape, text, size, color, bold=False, align="ctr", anchor="m", italic=False):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    p = tf.paragraphs[0]
    rtl_par(p, align)
    add_runs(p, text, size, color, bold, italic)


# ---------------------------------------------------------------- أشكال
def box(slide, x, y, w, h, fill=CARD, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.06, dash=False, line_w=1.0, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    s.shadow.inherit = False
    return s


def pill(slide, x, y, w, h, text, fill, color, size=14, bold=True, name=None):
    s = box(slide, x, y, w, h, fill=fill, line=None, radius=0.5, name=name)
    shape_text(s, text, size, color, bold)
    return s


def badge(slide, x, y, text, fill=PRIMARY, color=WHITE, size=16, name=None):
    """دائرة رقم أو كبسولة إن كان النص أطول من حرفين."""
    w = 0.55 if len(text) <= 2 else 0.85
    return pill(slide, x, y, w, 0.55, text, fill, color, size, True, name=name), w


def est_lines(text: str, width_in: float, size: float) -> int:
    clean = re.sub(r"\*\*", "", text)
    cap = max(4, int(width_in * 72 / (size * 0.56)))
    return max(1, math.ceil(len(clean) / cap))


# ---------------------------------------------------------------- هيكل الشريحة
class Ctx:
    def __init__(self, prs, meta, total):
        self.prs = prs
        self.meta = meta
        self.total = total


def new_slide(ctx: Ctx, dark=False):
    s = ctx.prs.slides.add_slide(ctx.prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(PRIMARY_DARK if dark else BG)
    return s


def add_slide_number(slide, x, y, w, h, color):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = "رقم الشريحة"
    tf = tb.text_frame
    p = tf.paragraphs[0]
    rtl_par(p, "l")
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
    fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"))
    rpr.set("lang", "en-US")
    rpr.set("sz", "1200")
    sf = etree.SubElement(rpr, qn("a:solidFill"))
    c = etree.SubElement(sf, qn("a:srgbClr"))
    c.set("val", color)
    lat = etree.SubElement(rpr, qn("a:latin"))
    lat.set("typeface", FONT)
    t = etree.SubElement(fld, qn("a:t"))
    t.text = "‹#›"
    # a:fld يجب أن يسبق a:endParaRPr إن وُجد
    end = p._p.find(qn("a:endParaRPr"))
    if end is not None:
        p._p.remove(end)
        p._p.append(end)


def footer(ctx: Ctx, slide, dark=False):
    meta = ctx.meta
    col = "A7F3D0" if dark else MUTED
    # شعار نصي: كبسولة SVU + اسم الجامعة (يمين)
    pill(slide, RIGHT - 0.75, 6.98, 0.75, 0.32, meta.get("logo_text", "SVU"),
         PRIMARY if not dark else CARD, WHITE if not dark else PRIMARY_DARK, 12, name="شعار الجامعة")
    textbox(slide, RIGHT - 0.75 - 6.2, 6.97, 6.1, 0.34,
            f"{meta['university']}  |  {meta['short_title']}", size=12, color=col,
            anchor="m", name="تذييل")
    add_slide_number(slide, MX, 6.97, 1.0, 0.34, col)


def title_bar(slide, text, dark=False):
    textbox(slide, MX, 0.38, CONTENT_W, 0.95, text, size=30, bold=True,
            color=WHITE if dark else PRIMARY_DARK, anchor="m", name="العنوان")


def status_band(slide, text, kind="pending"):
    """شريط حالة تحت العنوان: pending (كهرماني) أو badge (أخضر مزرق)."""
    if kind == "pending":
        fill, line, col = ACCENT_SOFT, ACCENT, ACCENT_INK
    else:
        fill, line, col = PRIMARY_SOFT, PRIMARY, PRIMARY_DARK
    b = box(slide, MX, 1.38, CONTENT_W, 0.46, fill=fill, line=line, radius=0.25,
            dash=(kind == "pending"), name="شريط الحالة")
    shape_text(b, text, 15, col, True, align="r")
    b.text_frame.margin_right = Inches(0.25)
    return 2.05


def top_of_content(spec, slide):
    if spec.get("pending"):
        return status_band(slide, spec["pending"], "pending")
    if spec.get("badge"):
        return status_band(slide, spec["badge"], "badge")
    return 1.6


def callout(slide, text, y, h=0.95, dark=False):
    b = box(slide, MX, y, CONTENT_W, h, fill=PRIMARY if not dark else PRIMARY_SOFT,
            line=None, radius=0.12, name="خلاصة")
    shape_text(b, text, 22, WHITE if not dark else PRIMARY_DARK, True, align="r")
    b.text_frame.margin_right = Inches(0.3)
    b.text_frame.margin_left = Inches(0.3)
    return b


def set_notes(slide, text, ctx):
    if not text:
        return
    tf = slide.notes_slide.notes_text_frame
    tf.text = text
    for p in tf.paragraphs:
        rtl_par(p, "r")
        for r in p.runs:
            rpr = r._r.get_or_add_rPr()
            rpr.set("lang", LANG)


# ---------------------------------------------------------------- قراءة جداول markdown
def read_md_table(path: Path, heading_contains: str):
    lines = path.read_text(encoding="utf-8").splitlines()
    start = next((i for i, l in enumerate(lines) if l.startswith("#") and heading_contains in l), None)
    if start is None:
        raise SystemExit(f"لم يوجد عنوان يحوي «{heading_contains}» في {path}")
    i = start + 1
    while i < len(lines) and not lines[i].lstrip().startswith("|"):
        i += 1
    rows = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows[0], rows[1:]


def first_float(s: str) -> float:
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        raise SystemExit(f"لا رقم في الخلية: {s!r}")
    return float(m.group())


# ---------------------------------------------------------------- أنواع الشرائح
def build_title(ctx, spec):
    s = new_slide(ctx, dark=True)
    # زخرفة: دوائر على اليسار
    for (x, y, d, c) in [(-2.2, 3.4, 6.2, PRIMARY_DEEP), (0.3, -1.6, 3.6, PRIMARY), (-0.6, 5.6, 2.6, PRIMARY)]:
        o = box(s, x, y, d, d, fill=c, line=None, shape=MSO_SHAPE.OVAL, name="زخرفة")
    # شعار نصي
    pill(s, RIGHT - 0.95, 0.55, 0.95, 0.42, ctx.meta.get("logo_text", "SVU"), CARD, PRIMARY_DARK, 16, name="شعار الجامعة")
    textbox(s, RIGHT - 0.95 - 5.6, 0.55, 5.5, 0.42, ctx.meta["university"], size=20, bold=True,
            color=WHITE, anchor="m", name="اسم الجامعة")
    pill(s, RIGHT - 3.8, 1.75, 3.8, 0.45, spec.get("subtitle", ""), ACCENT, WHITE, 16, name="نوع الرسالة")
    textbox(s, MX + 3.2, 2.35, RIGHT - MX - 3.2, 2.3, spec["title"], size=40, bold=True,
            color=WHITE, anchor="t", line=1.1, name="العنوان")
    textbox(s, MX + 3.2, 4.85, RIGHT - MX - 3.2, 1.6, spec.get("lines", []), size=22,
            color=PRIMARY_SOFT, line=1.2, after=4, name="بيانات الرسالة")
    set_notes(s, spec.get("notes"), ctx)
    # لا رقم/تذييل على شريحة العنوان
    return s


def build_section(ctx, spec):
    s = new_slide(ctx, dark=True)
    textbox(s, MX, 2.6, CONTENT_W, 1.2, spec["title"], size=44, bold=True, color=WHITE, anchor="m")
    if spec.get("subtitle"):
        textbox(s, MX, 3.9, CONTENT_W, 0.8, spec["subtitle"], size=24, color=PRIMARY_SOFT)
    footer(ctx, s, True)
    set_notes(s, spec.get("notes"), ctx)
    return s


def build_bullets(ctx, spec):
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    items = spec["items"]
    has_callout = bool(spec.get("callout"))
    bottom = (5.4 if has_callout else BOTTOM) - 0.1
    gap = 0.22
    n = len(items)
    rh = min(1.35, (bottom - y0 - gap * (n - 1)) / n)
    dot_w = 0.5
    for i, it in enumerate(items):
        y = y0 + i * (rh + gap)
        box(s, MX, y, CONTENT_W, rh, name="صف")
        badge(s, RIGHT - 0.25 - dot_w, y + (rh - 0.55) / 2, str(i + 1), name="رقم")
        if isinstance(it, dict):
            paras = [f"**{it['head']}**  {it['text']}"]
        else:
            paras = [it]
        textbox(s, MX + 0.25, y, CONTENT_W - 0.25 - dot_w - 0.55, rh, paras, size=24, anchor="m", line=1.1)
    if has_callout:
        callout(s, spec["callout"], 5.55, 1.15)
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


def build_cards(ctx, spec):
    dark = spec.get("theme") == "dark"
    s = new_slide(ctx, dark=dark)
    title_bar(s, spec["title"], dark)
    y0 = top_of_content(spec, s)
    if spec.get("lead"):
        textbox(s, MX, y0 - 0.05, CONTENT_W, 0.5, spec["lead"], size=20,
                color=PRIMARY_SOFT if dark else MUTED, anchor="m", bold=True)
        y0 += 0.6
    items = spec["items"]
    cols = spec.get("cols", 2)
    rows = math.ceil(len(items) / cols)
    has_callout = bool(spec.get("callout"))
    bottom = (5.45 if has_callout else BOTTOM)
    gx, gy = 0.3, 0.3
    cw = (CONTENT_W - gx * (cols - 1)) / cols
    big = cols <= 3
    ch = min(3.0 if big else 2.3, (bottom - y0 - gy * (rows - 1)) / rows)
    head_sz, text_sz = (24, 20) if big else (20, 16)
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = RIGHT - (c + 1) * cw - c * gx
        y = y0 + r * (ch + gy)
        box(s, x, y, cw, ch, fill=CARD, line=BORDER, name="بطاقة")
        if cols >= 5:                      # تخطيط مضغوط: رقم ثم عنوان تحته
            badge(s, x + cw - 0.2 - 0.55, y + 0.15, it["n"], fill=PRIMARY, size=15, name="رقم")
            textbox(s, x + 0.1, y + 0.8, cw - 0.2, ch - 0.9, it["head"], size=17, bold=True, anchor="t", line=1.05)
            continue
        bd, bw = badge(s, x + cw - 0.2 - (0.55 if len(it["n"]) <= 2 else 0.85), y + 0.2, it["n"],
                       fill=PRIMARY, size=16 if len(it["n"]) <= 2 else 14, name="رقم")
        hx = x + 0.15
        hw = cw - 0.15 - 0.2 - bw - 0.1
        textbox(s, hx, y + 0.12, hw, 0.75, it["head"], size=head_sz, bold=True, anchor="m", line=1.0)
        textbox(s, x + 0.15, y + 0.95, cw - 0.3, ch - 1.05, it["text"], size=text_sz,
                color=MUTED, anchor="t", line=1.15)
    if has_callout:
        callout(s, spec["callout"], min(5.6, y0 + rows * ch + (rows - 1) * gy + 0.4), 1.1)
    footer(ctx, s, dark)
    set_notes(s, spec.get("notes"), ctx)
    return s


def build_stats(ctx, spec):
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    items = spec["items"]
    cols = 3
    rows = math.ceil(len(items) / cols)
    gx, gy = 0.3, 0.3
    cap_h = 0.7 if spec.get("caption") else 0
    cw = (CONTENT_W - gx * (cols - 1)) / cols
    ch = (BOTTOM - cap_h - y0 - gy * (rows - 1)) / rows
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = RIGHT - (c + 1) * cw - c * gx
        y = y0 + r * (ch + gy)
        box(s, x, y, cw, ch, name="بطاقة رقم")
        textbox(s, x + 0.1, y + 0.18, cw - 0.2, ch * 0.45, LRM + it["value"] + LRM, size=40, bold=True,
                color=PRIMARY, align="ctr", anchor="m", name="رقم كبير")
        textbox(s, x + 0.2, y + 0.18 + ch * 0.45, cw - 0.4, ch * 0.55 - 0.3, it["label"], size=16,
                color=MUTED, align="ctr", anchor="t", line=1.1)
    if spec.get("caption"):
        textbox(s, MX, BOTTOM - cap_h + 0.08, CONTENT_W, cap_h - 0.1, spec["caption"], size=14,
                color=MUTED, anchor="m")
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


def build_flow(ctx, spec):
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    rows = spec["rows"]
    has_callout = bool(spec.get("callout"))
    bottom = 5.55 if has_callout else BOTTOM
    lab_h = 0.4
    gap = 0.25
    row_h = (bottom - y0 - gap * (len(rows) - 1)) / len(rows) - lab_h
    arrow_w = 0.42
    for ri, row in enumerate(rows):
        y = y0 + ri * (row_h + lab_h + gap)
        textbox(s, MX, y, CONTENT_W, lab_h, row["label"], size=17, bold=True, color=ACCENT_INK, anchor="m")
        steps = row["steps"]
        n = len(steps)
        sw = (CONTENT_W - (n - 1) * (arrow_w + 0.16)) / n
        for i, st in enumerate(steps):
            x = RIGHT - (i + 1) * sw - i * (arrow_w + 0.16)
            box(s, x, y + lab_h, sw, row_h, fill=CARD, line=PRIMARY, line_w=1.5, name="خطوة")
            textbox(s, x + 0.1, y + lab_h + 0.1, sw - 0.2, row_h * 0.42, st["head"], size=22, bold=True,
                    color=PRIMARY, align="ctr", anchor="m", line=1.0)
            textbox(s, x + 0.1, y + lab_h + row_h * 0.45, sw - 0.2, row_h * 0.52, st["text"], size=16,
                    color=MUTED, align="ctr", anchor="t", line=1.1)
            if i < n - 1:
                ax = x - 0.16 - arrow_w
                a = s.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(ax), Inches(y + lab_h + row_h / 2 - 0.2),
                                       Inches(arrow_w), Inches(0.4))
                a.name = "سهم"
                a.fill.solid()
                a.fill.fore_color.rgb = rgb(ACCENT)
                a.line.fill.background()
                a.shadow.inherit = False
    if has_callout:
        callout(s, spec["callout"], 5.7, 0.95)
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


def build_timeline(ctx, spec):
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    cuts = spec["cuts"]
    win = spec.get("window", 14)
    span = max(cuts) + win
    scale = CONTENT_W / span          # بوصة لكل يوم، اليوم 0 عند اليمين
    # مفتاح
    lx = RIGHT
    for lab, fill, col in [("تدريب", PRIMARY_SOFT, PRIMARY_DARK), ("تحقق (للضبط)", PRIMARY, WHITE),
                           ("اختبار (14 يوماً)", ACCENT, WHITE)]:
        w = 1.9
        lx -= w
        b = box(s, lx, y0, w, 0.36, fill=fill, line=PRIMARY if fill == PRIMARY_SOFT else None, radius=0.3, name="مفتاح")
        shape_text(b, lab, 14, col, True)
        lx -= 0.2
    row_y = y0 + 0.5
    pitch = 0.98
    for i, t in enumerate(cuts):
        y = row_y + i * pitch
        textbox(s, MX, y, CONTENT_W, 0.32, f"نقطة القطع t = {t} (اليوم من بداية العرض)", size=15, bold=True,
                color=PRIMARY_DARK, anchor="m")
        segs = [(0, t - win, PRIMARY_SOFT, PRIMARY_DARK, "تدريب"),
                (t - win, t, PRIMARY, WHITE, "تحقق"),
                (t, t + win, ACCENT, WHITE, "اختبار")]
        for a, b_, fill, col, lab in segs:
            x = RIGHT - b_ * scale
            w = (b_ - a) * scale
            r = box(s, x, y + 0.34, w, 0.5, fill=fill, line=PRIMARY if fill == PRIMARY_SOFT else None,
                    shape=MSO_SHAPE.RECTANGLE, name="مقطع")
            shape_text(r, lab, 13, col, True)
    chips = spec.get("chips", [])
    if chips:
        cy = row_y + len(cuts) * pitch + 0.15
        n = len(chips)
        g = 0.25
        cw = (CONTENT_W - g * (n - 1)) / n
        ch = BOTTOM - cy
        for i, c in enumerate(chips):
            x = RIGHT - (i + 1) * cw - i * g
            box(s, x, cy, cw, ch, name="بطاقة")
            textbox(s, x + 0.1, cy + 0.1, cw - 0.2, ch * 0.5, c["head"], size=18, bold=True, color=PRIMARY,
                    align="ctr", anchor="m", line=1.0)
            textbox(s, x + 0.1, cy + ch * 0.52, cw - 0.2, ch * 0.45, c["text"], size=14, color=MUTED,
                    align="ctr", anchor="t", line=1.05)
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


def items_column(s, x, y, w, h, heading, items, size=20):
    """بطاقة عمود بعنوان وصفوف نقطية."""
    box(s, x, y, w, h, name="عمود")
    hh = 0.75
    hb = box(s, x, y, w, hh, fill=PRIMARY, line=None, radius=0.2, name="عنوان العمود")
    shape_text(hb, heading, 20 if w > 5 else 18, WHITE, True, align="r")
    hb.text_frame.margin_right = Inches(0.2)
    n = len(items)
    avail = h - hh - 0.25
    rh = avail / n
    for i, it in enumerate(items):
        yy = y + hh + 0.15 + i * rh
        d = box(s, x + w - 0.42, yy + rh / 2 - 0.09, 0.18, 0.18, fill=ACCENT, line=None,
                shape=MSO_SHAPE.OVAL, name="نقطة")
        textbox(s, x + 0.15, yy, w - 0.15 - 0.55, rh, it, size=size, anchor="m", line=1.05)


def build_two_column(ctx, spec):
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    cols = spec["columns"]
    n = len(cols)
    g = 0.3
    cw = (CONTENT_W - g * (n - 1)) / n
    h = BOTTOM - y0
    size = 20 if n == 2 else 17
    for i, c in enumerate(cols):
        x = RIGHT - (i + 1) * cw - i * g
        items_column(s, x, y0, cw, h, c["heading"], c["items"], size=size)
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


def build_image(ctx, spec):
    from PIL import Image
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    path = ROOT / spec["image"]
    if not path.exists():
        raise SystemExit(f"الصورة غير موجودة: {path}")
    with Image.open(path) as im:
        iw, ih = im.size
    crop_b = spec.get("crop_bottom", 0.0)
    crop_t = spec.get("crop_top", 0.0)
    asp = iw / (ih * (1 - crop_b - crop_t))
    cap_h = 0.4 if spec.get("caption") else 0
    y0 = min(y0, 1.5)
    max_w, max_h = 7.8, 6.85 - y0 - cap_h
    pw = min(max_w, max_h * asp)
    ph = pw / asp
    px = MX                                # الصورة يساراً والنص يميناً (RTL)
    py = y0 + (max_h - ph) / 2
    pic = s.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    pic.name = "صورة"
    if crop_b:
        pic.crop_bottom = crop_b
    if crop_t:
        pic.crop_top = crop_t
    pic.line.color.rgb = rgb(BORDER)
    pic.line.width = Pt(1)
    pic._element.nvPicPr.cNvPr.set("descr", spec.get("alt") or spec.get("caption") or spec["title"])
    if spec.get("caption"):
        textbox(s, px, py + ph + 0.05, pw, cap_h, spec["caption"], size=13, color=MUTED, align="ctr", anchor="m")
    # عمود النص
    tx_w = CONTENT_W - pw - 0.4
    tx_x = RIGHT - tx_w
    items = spec.get("items", [])
    n = len(items)
    gap = 0.25
    rh = min(1.6, (BOTTOM - y0 - gap * (n - 1)) / n)
    for i, it in enumerate(items):
        y = y0 + i * (rh + gap)
        box(s, tx_x, y, tx_w, rh, name="بطاقة")
        badge(s, tx_x + tx_w - 0.2 - 0.55, y + 0.2, str(i + 1), name="رقم")
        textbox(s, tx_x + 0.15, y + 0.08, tx_w - 0.2 - 0.55 - 0.25, 0.65, it["head"], size=21, bold=True,
                anchor="m", line=1.0)
        textbox(s, tx_x + 0.15, y + 0.78, tx_w - 0.3, rh - 0.85, it["text"], size=17, color=MUTED, line=1.1)
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


def style_cell(cell, text, size, color, bold=False, fill=None, italic=False, align="r"):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    rtl_par(p, align)
    add_runs(p, text, size, color, bold, italic)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.1)
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)


def build_table(ctx, spec):
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    columns, rows = spec.get("columns"), spec.get("rows")
    if spec.get("source"):
        src = spec["source"]
        columns, rows = read_md_table(ROOT / src["file"], src["heading_contains"])
        if src.get("cols"):
            idx = src["cols"]
            columns = [columns[i] for i in idx]
            rows = [[r[i] for i in idx] for r in rows]
        if src.get("max_rows"):
            rows = rows[: src["max_rows"]]
    weights = spec.get("col_weights") or [1] * len(columns)
    cap_h = 0.55 if spec.get("caption") else 0
    n_rows = len(rows) + 1
    row_h = min(0.8, (BOTTOM - cap_h - y0) / n_rows)
    tbl_h = row_h * n_rows
    gs = s.shapes.add_table(n_rows, len(columns), Inches(MX), Inches(y0), Inches(CONTENT_W), Inches(tbl_h))
    gs.name = "جدول"
    tbl = gs.table
    # RTL: العمود الأول عند اليمين
    # LibreOffice يتجاهل rtl على الجدول، فيُعكس ترتيب الأعمدة يدوياً
    # ليظهر العمود الأول عند اليمين في كل التطبيقات
    tbl._tbl.tblPr.set("bandRow", "0")
    columns = list(reversed(columns))
    rows = [list(reversed(r)) for r in rows]
    weights = list(reversed(weights))
    tot = sum(weights)
    for i, wgt in enumerate(weights):
        tbl.columns[i].width = Emu(int(Inches(CONTENT_W) * wgt / tot))
    for r in range(n_rows):
        tbl.rows[r].height = Inches(row_h)
    size = spec.get("size", 16)
    for j, h in enumerate(columns):
        style_cell(tbl.cell(0, j), h, size, WHITE, True, PRIMARY)
    for i, row in enumerate(rows, start=1):
        fill = CARD if i % 2 else "EEF2F6"
        if (i - 1) in spec.get("explore_rows", []):          # صف استكشافي بتمييز واضح
            fill = PRIMARY_SOFT
            for j, v in enumerate(row):
                last = j == len(row) - 1
                style_cell(tbl.cell(i, j), v, size, ACCENT_INK if v.startswith("[") else PRIMARY_DARK,
                           bold=(j in (0, len(row) - 1)), fill=ACCENT_SOFT if v.startswith("[") else fill,
                           italic=v.startswith("["), align="ctr" if (v.startswith("[") or last) else "r")
            continue
        for j, v in enumerate(row):
            ph = v.strip().startswith("[")
            style_cell(tbl.cell(i, j), v, size, ACCENT_INK if ph else PRIMARY_DARK,
                       bold=(j == len(row) - 1), fill=ACCENT_SOFT if ph else fill, italic=ph,
                       align="ctr" if ph else "r")
    if spec.get("caption"):
        textbox(s, MX, y0 + tbl_h + 0.1, CONTENT_W, cap_h, spec["caption"], size=14, color=MUTED, anchor="m")
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


COLOR_NAMES = {"amber": ACCENT, "teal": PRIMARY, "gray": "94A3B8"}


def build_chart(ctx, spec):
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    ch = spec["chart"]
    if ch.get("source"):
        src = ch["source"]
        header, rows = read_md_table(ROOT / src["file"], src["heading_contains"])
        labels = [r[src.get("label_col", 0)] for r in rows]
        vals = [first_float(r[src.get("value_col", 1)]) for r in rows]
    else:
        labels, vals = ch["categories"], ch["values"]
    pairs = list(zip(labels, vals))
    if ch.get("sort") == "desc":
        pairs.sort(key=lambda p: -p[1])
    labels = [p[0] for p in pairs]
    vals = [p[1] for p in pairs]
    disp = [l.replace("_", " ") for l in labels]
    cd = CategoryChartData()
    cd.categories = disp
    cd.add_series(ch.get("series_name", "القيمة"), vals)
    cap_h = 0.45 if spec.get("caption") else 0
    side_w = 4.0 if spec.get("items") else 0
    cw = CONTENT_W - side_w - (0.3 if side_w else 0)
    cx = RIGHT - cw
    chh = BOTTOM - cap_h - y0
    chart_title_h = 0.5 if ch.get("title") else 0
    if ch.get("title"):
        textbox(s, cx, y0, cw, 0.45, ch["title"], size=16, bold=True, anchor="m", name="عنوان المخطط")
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(cx), Inches(y0 + chart_title_h), Inches(cw),
                            Inches(chh - chart_title_h), cd)
    gf.name = "مخطط أعمدة"
    chart = gf.chart
    chart.font.size = Pt(13)
    chart.font.name = FONT
    chart.font.color.rgb = rgb(PRIMARY_DARK)
    chart.has_legend = False
    chart.has_title = False          # العنوان مربع نص RTL فوق المخطط (LibreOffice لا يطبّق RTL على عنوان المخطط)
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = ch.get("number_format", "0.000")
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(13)
    dl.font.bold = True
    dl.font.name = FONT
    dl.font.color.rgb = rgb(PRIMARY_DARK)
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(PRIMARY)
    hl = ch.get("highlight", {})
    for i, lab in enumerate(labels):
        if lab in hl:
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(COLOR_NAMES.get(hl[lab], hl[lab]))
    ca = chart.category_axis
    ca.reverse_order = True
    ca.tick_labels.font.size = Pt(13)
    ca.format.line.color.rgb = rgb(BORDER)
    ca.has_major_gridlines = False
    va = chart.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale = 0
    va.maximum_scale = max(vals) * 1.18
    # بطاقات جانبية
    items = spec.get("items", [])
    if items:
        sx = MX
        n = len(items)
        g = 0.25
        ih = (BOTTOM - cap_h - y0 - g * (n - 1)) / n
        for i, it in enumerate(items):
            y = y0 + i * (ih + g)
            box(s, sx, y, side_w, ih, name="بطاقة")
            textbox(s, sx + 0.1, y + 0.05, side_w - 0.2, ih * 0.38, LRM + it["value"] + LRM, size=32, bold=True,
                    color=ACCENT if i == 0 else PRIMARY, align="ctr", anchor="m")
            textbox(s, sx + 0.15, y + ih * 0.42, side_w - 0.3, ih * 0.56, it["label"], size=14, color=MUTED,
                    align="ctr", anchor="t", line=1.05)
    if spec.get("caption"):
        textbox(s, MX, BOTTOM - cap_h + 0.05, CONTENT_W, cap_h, spec["caption"], size=13, color=MUTED, anchor="m")
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


def build_quote(ctx, spec):
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    textbox(s, MX + 0.6, y0, 1.2, 1.4, "”", size=96, bold=True, color=ACCENT, align="r")
    textbox(s, MX + 0.6, y0 + 0.9, CONTENT_W - 1.2, 2.6, spec["text"], size=32, italic=True, anchor="t", line=1.2)
    if spec.get("by"):
        textbox(s, MX + 0.6, y0 + 3.7, CONTENT_W - 1.2, 0.5, "— " + spec["by"], size=18, color=MUTED)
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


def build_placeholder(ctx, spec):
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    panels = spec["panels"]
    n = len(panels)
    cols = n if n <= 4 else 2
    rows = math.ceil(n / cols)
    has_callout = bool(spec.get("callout"))
    bottom = 5.6 if has_callout else BOTTOM
    g = 0.3
    cw = (CONTENT_W - g * (cols - 1)) / cols
    ch = (bottom - y0 - g * (rows - 1)) / rows
    for i, p in enumerate(panels):
        r, c = divmod(i, cols)
        x = RIGHT - (c + 1) * cw - c * g
        y = y0 + r * (ch + g)
        box(s, x, y, cw, ch, name="لوحة")
        textbox(s, x + 0.15, y + 0.1, cw - 0.3, 0.75, p["head"], size=20 if cols <= 3 else 18, bold=True, anchor="m", line=1.0)
        textbox(s, x + 0.15, y + 0.9, cw - 0.3, 0.9, p["text"], size=15, color=MUTED, line=1.05)
        top = y + 1.9
        img = ROOT / p["image"] if p.get("image") else None
        if img is not None and img.exists():     # شكل جاهز (PNG) يحلّ محلّ المكان النائب
            from PIL import Image
            with Image.open(img) as im:
                iw, ih = im.size
            aw, ah = cw - 0.4, y + ch - 0.2 - top
            pw = min(aw, ah * iw / ih)
            ph = pw * ih / iw
            pic = s.shapes.add_picture(str(img), Inches(x + 0.2 + (aw - pw) / 2), Inches(top + (ah - ph) / 2),
                                       Inches(pw), Inches(ph))
            pic.name = "شكل"
            pic._element.nvPicPr.cNvPr.set("descr", p.get("alt") or p["head"])
        else:
            slot = box(s, x + 0.2, top, cw - 0.4, y + ch - 0.2 - top, fill=ACCENT_SOFT, line=ACCENT,
                       dash=True, radius=0.08, line_w=1.5, name="مكان النتيجة")
            shape_text(slot, p["slot"], 18, ACCENT_INK, True, italic=True)
    if has_callout:
        callout(s, spec["callout"], 5.75, 0.95)
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


def build_qa(ctx, spec):
    if spec.get("variant") == "closing":
        s = new_slide(ctx, dark=True)
        for (x, y, d, c) in [(-2.2, 3.4, 6.2, PRIMARY_DEEP), (0.3, -1.6, 3.6, PRIMARY)]:
            box(s, x, y, d, d, fill=c, line=None, shape=MSO_SHAPE.OVAL, name="زخرفة")
        textbox(s, MX + 3.2, 1.7, RIGHT - MX - 3.2, 1.4, spec["title"], size=54, bold=True, color=WHITE,
                anchor="m", name="العنوان")
        pill(s, RIGHT - 3.6, 3.3, 3.6, 0.6, spec.get("subtitle", ""), ACCENT, WHITE, 24, name="أسئلة")
        textbox(s, MX + 3.2, 4.5, RIGHT - MX - 3.2, 1.4, spec.get("lines", []), size=22, color=PRIMARY_SOFT,
                after=4, line=1.2)
        footer(ctx, s, True)
        set_notes(s, spec.get("notes"), ctx)
        return s
    s = new_slide(ctx)
    title_bar(s, spec["title"])
    y0 = top_of_content(spec, s)
    qs = spec["questions"]
    n = len(qs)
    g = 0.3
    cw = (CONTENT_W - g * (n - 1)) / n
    h = BOTTOM - y0
    for i, q in enumerate(qs):
        x = RIGHT - (i + 1) * cw - i * g
        box(s, x, y0, cw, h, name="بطاقة")
        qb = box(s, x, y0, cw, 0.9, fill=PRIMARY_SOFT, line=None, radius=0.2, name="السؤال")
        shape_text(qb, "س: " + q["q"], 24, PRIMARY_DARK, True, align="r")
        qb.text_frame.margin_right = Inches(0.25)
        ans = q["a"]
        m = len(ans)
        avail = h - 0.9 - 0.3
        size = 20 if n == 1 else 17
        rh = avail / m
        for k, a in enumerate(ans):
            yy = y0 + 0.9 + 0.15 + k * rh
            box(s, x + cw - 0.42, yy + rh / 2 - 0.09, 0.18, 0.18, fill=ACCENT, line=None,
                shape=MSO_SHAPE.OVAL, name="نقطة")
            textbox(s, x + 0.15, yy, cw - 0.15 - 0.55, rh, a, size=size, anchor="m", line=1.05)
    footer(ctx, s)
    set_notes(s, spec.get("notes"), ctx)
    return s


BUILDERS = {
    "title": build_title, "section": build_section, "bullets": build_bullets, "cards": build_cards,
    "stats": build_stats, "flow": build_flow, "timeline": build_timeline, "two-column": build_two_column,
    "image": build_image, "table": build_table, "chart": build_chart, "quote": build_quote,
    "qa": build_qa, "placeholder": build_placeholder,
}


# ---------------------------------------------------------------- الدخول
def load_spec(path: Path) -> dict:
    if path.suffix in (".yaml", ".yml"):
        import yaml  # noqa
        return yaml.safe_load(path.read_text(encoding="utf-8"))
    return json.loads(path.read_text(encoding="utf-8"))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", help="مسار مخطط الشرائح (yaml/json)")
    ap.add_argument("--out", help="مسار ملف pptx الناتج")
    a = ap.parse_args(argv)
    if a.spec:
        spec_path = Path(a.spec)
    else:
        y = ROOT / "docs/working/07-presentation.yaml"
        spec_path = y if y.exists() else ROOT / "docs/working/07-presentation.json"
    spec = load_spec(spec_path)
    meta = spec["meta"]
    out = Path(a.out) if a.out else ROOT / meta["output"]
    out.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(W_IN)
    prs.slide_height = Inches(H_IN)
    cp = prs.core_properties
    cp.title = meta["thesis_title"]
    cp.author = meta["author"]
    cp.subject = f"{meta['short_title']} — {meta['university']}"
    cp.language = LANG
    slides = spec["slides"]
    ctx = Ctx(prs, meta, len(slides))
    for i, sp in enumerate(slides, start=1):
        kind = sp["type"]
        if kind not in BUILDERS:
            raise SystemExit(f"الشريحة {i}: نوع غير معروف {kind!r}")
        BUILDERS[kind](ctx, sp)
    prs.save(out)
    print(f"تم: {out} ({len(slides)} شريحة)")


if __name__ == "__main__":
    sys.exit(main())
