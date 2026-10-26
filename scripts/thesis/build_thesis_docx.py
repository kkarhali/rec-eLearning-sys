#!/usr/bin/env python3
"""Build the final thesis Word file: deliverables/06-final-thesis/الأطروحة-النهائية.docx

Pipeline
  1. Copy the Markdown sources of docs/working/06-thesis/ (front matter, chapters 1-7,
     references, appendices A-D) into a temporary folder and prepare them for the build:
       * front matter: drop the "# المقدمات" wrapper and the Arabic title page (the cover of
         build_docx.py replaces it), keep the English title page as plain centred paragraphs,
         promote its sections to level-1 headings, replace the placeholder section of the indexes
         with a marker, drop the horizontal rules;
       * figures: an image followed by its own italic caption line loses its alt text (no duplicate
         caption); the caption line is split into the caption proper and its note; every image gets a
         width that fits the text block (height capped);
  2. Run scripts/build_docx.py (pandoc + RTL post-processing) with the cover values and no TOC.
  3. Post-process with python-docx:
       * a page break before every level-1 heading (chapters, appendices, references, front sections);
       * the marker becomes three fields: table of contents (headings 1-3), list of figures and list
         of tables (built from the caption styles), with updateFields so Word fills them on open;
       * caption styles: figure captions "Image Caption", table titles "Table Caption", centred,
         kept with their figure or table;
       * the English title page and the English abstract are laid out left-to-right.

Equations (OMML) are kept as they are; this LibreOffice build has no Math module, so the fields are
not pre-rendered here (a LibreOffice round trip would drop the equations). Use --pdf to also write a
PDF preview (from a separate copy whose fields are refreshed by LibreOffice; equations are missing
in that preview only).

Usage (from the repository root):  python3 scripts/thesis/build_thesis_docx.py [--pdf[=DIR]]
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs" / "working" / "06-thesis"
OUT = ROOT / "deliverables" / "06-final-thesis" / "الأطروحة-النهائية.docx"
ORDER = ["00-front-matter.md", "ch1-introduction.md", "ch2-background-literature.md", "ch3-methodology.md",
         "ch4-analysis-design.md", "ch5-implementation-results.md", "ch6-discussion.md", "ch7-conclusion.md",
         "references.md", "appendix-A-comparison-matrix.md", "appendix-B-numeric-examples.md",
         "appendix-C-requirements.md", "appendix-D-confirmatory-details.md"]
TEXT_WIDTH_CM, MAX_HEIGHT_CM = 15.0, 18.0
MARKER = "@@INDEXES@@"
EN_TITLE = "A Recommendation System for E-Learning Based on User Behavior Analysis"
COVER = dict(
    title="نظام توصية للتعليم الإلكتروني القائم على تحليل سلوك المستخدم",
    title_en=EN_TITLE,
    subtitle="رسالة أُعدّت لنيل درجة الماجستير في علوم الوب",
    author="كنان محمد قرحالي",
    supervisor="د. باسل الخطيب",
    university="الجمهورية العربية السورية",
    faculty="الجامعة الافتراضية السورية",
    program="برنامج ماجستير علوم الوب",
    year="2026",
)
IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)\)\s*$")
FIGCAP_RE = re.compile(r"^\*(الشكل [0-9ء-ي]+-\d+: .*)\*\s*$")


# --------------------------------------------------------------------------- 1. prepare sources
def prepare_front(text: str) -> str:
    text = text.replace("# المقدمات\n", "", 1)
    a = text.index("## صفحة العنوان (بالعربية)")
    b = text.index("## Title Page (English)")
    text = text[:a] + text[b:]
    # English title page: plain paragraphs (no headings), the title marked for resizing
    a = text.index("## Title Page (English)")
    b = text.index("## الشكر والتقدير")
    block = text[a:b].replace("## Title Page (English)\n", "")
    block = block.replace(f"# {EN_TITLE}", f"**{EN_TITLE}**")
    block = block.replace("&nbsp;", "")
    lines = [l for l in block.split("\n") if l.strip() not in ("---",)]
    # one paragraph per line
    block = "\n\n".join(l.strip() for l in lines if l.strip())
    text = text[:a] + block + "\n\n" + text[b:]
    # indexes placeholder -> marker
    a = text.index("## فهرس المحتويات، وفهرس الجداول، وفهرس الأشكال")
    b = text.index("## قائمة المختصرات")
    text = text[:a] + MARKER + "\n\n" + text[b:]
    # front sections become level-1 headings
    text = re.sub(r"^## ", "# ", text, flags=re.M)
    text = re.sub(r"^---\s*$", "", text, flags=re.M)
    return text


def image_size(path: str) -> str:
    w, h = Image.open(ROOT / path).size
    if TEXT_WIDTH_CM * h / w > MAX_HEIGHT_CM:
        return f"{{height={MAX_HEIGHT_CM:g}cm}}"
    return f"{{width={TEXT_WIDTH_CM:g}cm}}"


def prepare_figures(text: str) -> tuple[str, list[str]]:
    lines = text.split("\n")
    missing = []
    for i, line in enumerate(lines):
        m = IMG_RE.match(line)
        if not m:
            continue
        alt, path = m.group(1), m.group(2)
        if not (ROOT / path).exists():
            missing.append(path)
            continue
        # own italic caption below (after a blank line)?
        j = i + 1
        while j < len(lines) and not lines[j].strip():
            j += 1
        own = j < len(lines) and FIGCAP_RE.match(lines[j])
        if own:
            cap = FIGCAP_RE.match(lines[j]).group(1)
            k = cap.find(". ")
            title, note = (cap[:k], cap[k + 2:]) if k > 0 else (cap.rstrip("."), "")
            # image without alt (no automatic caption) + caption paragraph + note paragraph
            lines[i] = f"![]({path}){image_size(path)}"
            lines[j] = f"::: {{custom-style=\"Image Caption\"}}\n{title}\n:::" + (f"\n\n::: {{custom-style=\"Caption Note\"}}\n{note}\n:::" if note else "")
        else:
            lines[i] = f"![{alt}]({path}){image_size(path)}"
    return "\n".join(lines), missing


def prepare(tmp: Path) -> list[Path]:
    out, missing = [], []
    for name in ORDER:
        text = (SRC / name).read_text(encoding="utf-8")
        if name == "00-front-matter.md":
            text = prepare_front(text)
        text, miss = prepare_figures(text)
        missing += miss
        p = tmp / name
        p.write_text(text, encoding="utf-8")
        out.append(p)
    if missing:
        sys.exit(f"missing images: {missing}")
    return out


# --------------------------------------------------------------------------- 3. post-process
def field_paragraph(doc, anchor, instr: str):
    """Insert a paragraph holding a complex field (TOC) before anchor."""
    p = OxmlElement("w:p")
    anchor.addprevious(p)
    ppr = OxmlElement("w:pPr")
    bidi = OxmlElement("w:bidi"); bidi.set(qn("w:val"), "1"); ppr.append(bidi)
    p.append(ppr)
    def run(child):
        r = OxmlElement("w:r"); r.append(child); p.append(r)
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); fc.set(qn("w:dirty"), "true"); run(fc)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; run(it)
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); run(fc)
    t = OxmlElement("w:t"); t.text = "يُحدَّث هذا الفهرس عند فتح الملف (تحديث الحقول)."; run(t)
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); run(fc)
    return p


def heading_paragraph(anchor, text: str):
    from docx.text.paragraph import Paragraph
    p = OxmlElement("w:p")
    anchor.addprevious(p)
    ppr = OxmlElement("w:pPr")
    ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), "TOCHeading"); ppr.append(ps)
    pb = OxmlElement("w:pageBreakBefore"); ppr.append(pb)
    bidi = OxmlElement("w:bidi"); bidi.set(qn("w:val"), "1"); ppr.append(bidi)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); ppr.append(jc)
    p.append(ppr)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    for tag in ("w:b", "w:bCs"):
        e = OxmlElement(tag); rpr.append(e)
    for tag in ("w:sz", "w:szCs"):
        e = OxmlElement(tag); e.set(qn("w:val"), "40"); rpr.append(e)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "1F3864"); rpr.append(col)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; r.append(t); p.append(r)
    return p


def set_ppr(p, tag, val=None):
    ppr = p._p.get_or_add_pPr()
    el = ppr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        ppr.append(el)
    if val is not None:
        el.set(qn("w:val"), val)
    return el


def ltr(p, align=WD_ALIGN_PARAGRAPH.LEFT):
    set_ppr(p, "w:bidi", "0")
    p.alignment = align
    for r in p.runs:
        rpr = r._r.get_or_add_rPr()
        rtl = rpr.find(qn("w:rtl"))
        if rtl is not None:
            rpr.remove(rtl)


def ensure_style(doc, name, base="Normal", size=None, italic=False, center=True, keep_next=False):
    styles = doc.styles
    st = next((x for x in styles if x.name == name or x.style_id == name.replace(" ", "")), None)
    if st is None:
        from docx.enum.style import WD_STYLE_TYPE
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = next(x for x in styles if x.style_id == base.replace(" ", ""))
    if size:
        st.font.size = Pt(size)
        rpr = st.element.get_or_add_rPr()
        szcs = rpr.find(qn("w:szCs"))
        if szcs is None:
            szcs = OxmlElement("w:szCs"); rpr.append(szcs)
        szcs.set(qn("w:val"), str(size * 2))
    st.font.italic = italic
    ppr = st.element.get_or_add_pPr()
    for tag, val in (("w:bidi", "1"), ("w:jc", "center" if center else "both")):
        el = ppr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); ppr.append(el)
        el.set(qn("w:val"), val)
    if keep_next and ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))
    return st


TEXT_TWIPS = int(15.5 / 2.54 * 1440)  # text block width (A4, margins 2.5 + 3.0 cm)


def fit_tables(doc) -> None:
    """Column widths from content (pandoc gives equal widths to long pipe tables) and smaller table text."""
    for table in doc.tables:
        rows = table.rows
        ncol = max(len(r.cells) for r in rows)
        longest_tok = [1] * ncol
        total_len = [0] * ncol
        for ri, r in enumerate(rows):
            for c, cell in enumerate(r.cells[:ncol]):
                txt = cell.text.strip()
                toks = txt.split() or [""]
                if ri > 0 or len(rows) == 1:  # header words may wrap; body words should not
                    # paths and identifiers break at / _ . so they need not fit on one line
                    lens = [min(len(t), 14) if re.search(r"[/_.]", t) and len(t) > 14 else len(t) for t in toks]
                    longest_tok[c] = max(longest_tok[c], min(max(lens), 22))
                    total_len[c] += len(txt)
        pt = 8 if ncol >= 10 else 9 if ncol >= 7 else 10 if ncol >= 5 else 11 if ncol >= 4 else 12
        char_tw = pt * 20 * 0.55  # average glyph width in twips
        need = [max(lt * char_tw + 160, 500) for lt in longest_tok]   # no word broken
        nbody = max(len(rows) - 1, 1)
        want = [max(n, min(tl / nbody, 60) * char_tw * 0.9) for n, tl in zip(need, total_len)]
        scale = TEXT_TWIPS / sum(want)
        widths = [int(w * scale) for w in want]
        if sum(need) <= TEXT_TWIPS and scale < 1:
            # shrink only the columns that have room above their minimum
            room = [w - n for w, n in zip(want, need)]
            excess = sum(want) - TEXT_TWIPS
            widths = [int(w - excess * r / sum(room)) if sum(room) else int(w) for w, r in zip(want, room)]
        tbl = table._tbl
        tblpr = tbl.tblPr
        tw = tblpr.find(qn("w:tblW"))
        if tw is None:
            tw = OxmlElement("w:tblW"); tblpr.append(tw)
        tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(sum(widths)))
        lay = tblpr.find(qn("w:tblLayout"))
        if lay is None:
            lay = OxmlElement("w:tblLayout"); tblpr.append(lay)
        lay.set(qn("w:type"), "fixed")
        mar = tblpr.find(qn("w:tblCellMar"))
        if mar is None:
            mar = OxmlElement("w:tblCellMar"); tblpr.append(mar)
        for side in ("left", "right", "start", "end"):
            e = mar.find(qn(f"w:{side}"))
            if e is None:
                e = OxmlElement(f"w:{side}"); mar.append(e)
            e.set(qn("w:w"), "57"); e.set(qn("w:type"), "dxa")
        grid = tbl.find(qn("w:tblGrid"))
        if grid is not None:
            for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
                gc.set(qn("w:w"), str(w))
        for r in rows:
            for c, cell in enumerate(r.cells[:ncol]):
                tcpr = cell._tc.get_or_add_tcPr()
                tcw = tcpr.find(qn("w:tcW"))
                if tcw is None:
                    tcw = OxmlElement("w:tcW"); tcpr.insert(0, tcw)
                tcw.set(qn("w:type"), "dxa"); tcw.set(qn("w:w"), str(widths[c]))
                for para in cell.paragraphs:
                    sp = set_ppr(para, "w:spacing"); sp.set(qn("w:line"), "240"); sp.set(qn("w:lineRule"), "auto")
                    sp.set(qn("w:after"), "0"); sp.set(qn("w:before"), "0")
                    for run in para.runs:
                        rpr = run._r.get_or_add_rPr()
                        for tag in ("w:sz", "w:szCs"):
                            e = rpr.find(qn(tag))
                            if e is None:
                                e = OxmlElement(tag); rpr.append(e)
                            e.set(qn("w:val"), str(pt * 2))
        # header row repeats on each page
        trpr = rows[0]._tr.get_or_add_trPr()
        if trpr.find(qn("w:tblHeader")) is None:
            trpr.append(OxmlElement("w:tblHeader"))


def postprocess(path: Path) -> dict:
    doc = Document(str(path))
    report = {"tables": 0, "figures": 0, "images": 0}
    # page break before every level-1 heading
    h1 = next(st for st in doc.styles if st.style_id == "Heading1").element.get_or_add_pPr()
    if h1.find(qn("w:pageBreakBefore")) is None:
        h1.append(OxmlElement("w:pageBreakBefore"))
    cap_img = ensure_style(doc, "Image Caption", size=12, center=True)
    cap_tab = ensure_style(doc, "Table Caption", size=12, center=True, keep_next=True)
    sp = cap_tab.element.get_or_add_pPr()
    spc = sp.find(qn("w:spacing"))
    if spc is None:
        spc = OxmlElement("w:spacing"); sp.append(spc)
    spc.set(qn("w:before"), "240"); spc.set(qn("w:after"), "120")
    ensure_style(doc, "Caption Note", size=11, italic=True, center=True)
    for st in (cap_img, cap_tab):
        st.font.bold = True
    paras = doc.paragraphs
    # English title page: from the paragraph after the cover break up to the first Heading 1
    first_h1 = next(i for i, p in enumerate(paras) if p.style.name == "Heading 1")
    for p in paras[:first_h1]:
        if (any(ch.isascii() and ch.isalpha() for ch in p.text) or p.text.strip().isdigit()) and not re.search(r"[؀-ۿ]", p.text):
            ltr(p, WD_ALIGN_PARAGRAPH.CENTER)
            if p.text.strip() == EN_TITLE:
                for r in p.runs:
                    r.font.size = Pt(20); r.bold = True
    # break before the English title page (after the Arabic cover there is already one) and after it:
    # the next Heading 1 has pageBreakBefore, so only the start needs care: the cover adds its own break.
    # English abstract: Heading "Abstract" up to the next Heading 1
    in_en = False
    for p in paras:
        if p.style.name == "Heading 1":
            in_en = p.text.strip() == "Abstract"
            if in_en:
                ltr(p, WD_ALIGN_PARAGRAPH.LEFT)
            continue
        if in_en:
            ltr(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
    # captions
    for p in doc.paragraphs:
        txt = p.text.strip()
        if re.match(r"^جدول [0-9ء-ي]+-\d+", txt) and p.runs and all(r.bold for r in p.runs if r.text.strip()):
            p.style = cap_tab
            report["tables"] += 1
        elif p.style.name == "Image Caption" or (p.style.name in ("Caption",) and txt.startswith("الشكل")):
            p.style = cap_img
            report["figures"] += 1
        if p._p.find(f".//{qn('w:drawing')}") is not None:
            report["images"] += 1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_ppr(p, "w:keepNext")
    fit_tables(doc)
    # indexes
    marker = next(p for p in doc.paragraphs if p.text.strip() == MARKER)
    anchor = marker._p
    heading_paragraph(anchor, "فهرس المحتويات")
    field_paragraph(doc, anchor, 'TOC \\o "1-3" \\h \\z \\u')
    heading_paragraph(anchor, "فهرس الأشكال")
    field_paragraph(doc, anchor, 'TOC \\h \\z \\t "Image Caption,1"')
    heading_paragraph(anchor, "فهرس الجداول")
    field_paragraph(doc, anchor, 'TOC \\h \\z \\t "Table Caption,1"')
    anchor.getparent().remove(anchor)
    settings = doc.settings.element
    uf = settings.find(qn("w:updateFields"))
    if uf is None:
        uf = OxmlElement("w:updateFields"); settings.append(uf)
    uf.set(qn("w:val"), "true")
    doc.save(str(path))
    return report


# --------------------------------------------------------------------------- main
def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        files = prepare(tmp)
        cmd = [sys.executable, str(ROOT / "scripts" / "build_docx.py"), str(OUT), *map(str, files), "--no-toc"]
        for k, v in COVER.items():
            cmd += [f"--{k.replace('_', '-')}", v]
        subprocess.run(cmd, check=True, cwd=ROOT)
    rep = postprocess(OUT)
    print(f"post-processed: {rep}")
    pdf = [a for a in sys.argv[1:] if a.startswith("--pdf")]
    if pdf:
        pdf_preview(Path(pdf[0].split("=", 1)[1]) if "=" in pdf[0] else OUT.parent / "preview")


def pdf_preview(prev: Path) -> None:
    """PDF preview from a copy whose index fields are refreshed by LibreOffice (equations missing there)."""
    prev.mkdir(parents=True, exist_ok=True)
    copy = prev / "thesis-preview.docx"
    shutil.copy(OUT, copy)
    for _ in range(2):  # twice: the second pass sees the indexes at their full length
        subprocess.run([sys.executable, str(ROOT / "scripts" / "lo_update_toc.py"), str(copy)], check=True, timeout=900)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(prev), str(copy)],
                   check=True, timeout=900, capture_output=True)
    print(f"preview: {prev / 'thesis-preview.pdf'}")


if __name__ == "__main__":
    main()
